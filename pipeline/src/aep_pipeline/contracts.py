"""Input contracts: declared fields, types, geometry, reference system and record counts.

Sensitivity classes
-------------------
public        may be written to curated artifacts
internal      may be read to link or tokenize, but is never written raw
drop          sensitive or unneeded, never read
undocumented  meaning is not documented, never read
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pyogrio
import yaml

from .errors import ConfigError, ContractError

SENSITIVITIES = ("public", "internal", "drop", "undocumented")
LFS_POINTER_PREFIX = b"version https://git-lfs"
SHAPEFILE_PARTS = (".shp", ".shx", ".dbf")


@dataclass(frozen=True)
class FieldSpec:
    name: str
    dtype: str
    sensitivity: str
    description: str
    placeholder_dates: bool = False


@dataclass(frozen=True)
class LayerContract:
    name: str
    path: str
    geometry_type: str
    min_records: int
    geometry_use: str
    crs: str | None = None
    crs_contains: str | None = None
    fields: dict[str, FieldSpec] = field(default_factory=dict)

    def columns(self, *sensitivities: str) -> list[str]:
        wanted = sensitivities or ("public", "internal")
        return [f.name for f in self.fields.values() if f.sensitivity in wanted]


@dataclass(frozen=True)
class TopologyContract:
    name: str
    path: str
    min_records: int
    premise_pattern: str
    field_delimiter: str
    node_delimiter: str
    node_parts: tuple[str, ...]


@dataclass(frozen=True)
class WorkbookContract:
    name: str
    path: str
    sheets: tuple[str, ...]
    min_rows: dict[str, int]


@dataclass(frozen=True)
class LayerCheck:
    name: str
    records: int
    unlisted_fields: tuple[str, ...]


def default_config_dir() -> Path:
    """The ``pipeline/config`` directory next to the package source."""
    return Path(__file__).resolve().parents[2] / "config"


def _read_yaml(path: Path) -> dict:
    try:
        with path.open("r", encoding="utf-8") as handle:
            data = yaml.safe_load(handle)
    except FileNotFoundError as error:
        raise ConfigError(f"Configuration file not found: {path}") from error
    if not isinstance(data, dict):
        raise ConfigError(f"Configuration file is not a mapping: {path}")
    return data


def _parse_layer(path: Path) -> LayerContract:
    data = _read_yaml(path)
    fields: dict[str, FieldSpec] = {}
    for fname, spec in (data.get("fields") or {}).items():
        sensitivity = spec.get("sensitivity")
        if sensitivity not in SENSITIVITIES:
            raise ConfigError(f"{path.name}: field '{fname}' has invalid sensitivity '{sensitivity}'")
        fields[fname] = FieldSpec(
            fname,
            spec["type"],
            sensitivity,
            spec.get("description", ""),
            bool(spec.get("placeholder_dates", False)),
        )
    return LayerContract(
        name=data["name"],
        path=data["path"],
        geometry_type=data["geometry_type"],
        min_records=int(data["min_records"]),
        geometry_use=data.get("geometry_use", "output"),
        crs=data.get("crs"),
        crs_contains=data.get("crs_contains"),
        fields=fields,
    )


def load_layer_contracts(config_dir: Path | None = None) -> dict[str, LayerContract]:
    """Load every shapefile contract in ``config/contracts``, keyed by layer name."""
    directory = (config_dir or default_config_dir()) / "contracts"
    contracts: dict[str, LayerContract] = {}
    for path in sorted(directory.glob("*.yaml")):
        data = _read_yaml(path)
        if data.get("kind") == "shapefile":
            contract = _parse_layer(path)
            contracts[contract.name] = contract
    if not contracts:
        raise ConfigError(f"No shapefile contracts found in {directory}")
    return contracts


def load_topology_contract(config_dir: Path | None = None) -> TopologyContract:
    data = _read_yaml((config_dir or default_config_dir()) / "contracts" / "matched_upstream.yaml")
    return TopologyContract(
        name=data["name"],
        path=data["path"],
        min_records=int(data["min_records"]),
        premise_pattern=data["premise_pattern"],
        field_delimiter=data["field_delimiter"],
        node_delimiter=data["node_delimiter"],
        node_parts=tuple(data["node_parts"]),
    )


def load_workbook_contract(config_dir: Path | None = None) -> WorkbookContract:
    data = _read_yaml((config_dir or default_config_dir()) / "contracts" / "Install_Labor_HackOHIO.yaml")
    return WorkbookContract(
        name=data["name"],
        path=data["path"],
        sheets=tuple(data["sheets"]),
        min_rows=dict(data.get("min_rows", {})),
    )


def is_lfs_pointer(path: Path) -> bool:
    """True when the file is a Git LFS pointer instead of real content."""
    try:
        with path.open("rb") as handle:
            return handle.read(len(LFS_POINTER_PREFIX)) == LFS_POINTER_PREFIX
    except OSError:
        return False


def _require_file(raw_dir: Path, relative: str, label: str) -> Path:
    path = raw_dir / relative
    if not path.is_file():
        raise ContractError(f"{label}: file not found at {relative}")
    if is_lfs_pointer(path):
        raise ContractError(
            f"{relative}: this is a Git LFS pointer file, not real data. "
            "Run 'git lfs pull' to fetch the data files."
        )
    return path


def validate_layer(raw_dir: Path, contract: LayerContract) -> LayerCheck:
    """Check one shapefile layer against its contract using file metadata only."""
    shp = _require_file(raw_dir, contract.path, contract.name)
    for suffix in SHAPEFILE_PARTS:
        part = shp.with_suffix(suffix)
        if not part.is_file():
            raise ContractError(f"{contract.path}: required part {part.name} is missing")
        if is_lfs_pointer(part):
            raise ContractError(
                f"{part.name}: this is a Git LFS pointer file, not real data. "
                "Run 'git lfs pull' to fetch the data files."
            )

    info = pyogrio.read_info(shp)

    if info["geometry_type"] != contract.geometry_type:
        raise ContractError(
            f"{contract.path}: geometry type is '{info['geometry_type']}', expected '{contract.geometry_type}'"
        )

    crs_text = info["crs"] or ""
    if contract.crs and crs_text != contract.crs:
        raise ContractError(f"{contract.path}: reference system is '{crs_text[:60]}', expected '{contract.crs}'")
    if contract.crs_contains and contract.crs_contains not in crs_text:
        raise ContractError(
            f"{contract.path}: reference system does not contain '{contract.crs_contains}'"
        )

    actual = dict(zip(info["fields"], (str(d) for d in info["dtypes"])))
    for fname, spec in contract.fields.items():
        if fname not in actual:
            raise ContractError(f"{contract.path}: field '{fname}' is missing")
        if actual[fname] != spec.dtype:
            raise ContractError(
                f"{contract.path}: field '{fname}' has type '{actual[fname]}', expected '{spec.dtype}'"
            )

    records = int(info["features"])
    if records < contract.min_records:
        raise ContractError(
            f"{contract.path}: has {records} records, expected at least {contract.min_records}"
        )

    unlisted = tuple(sorted(set(actual) - set(contract.fields)))
    return LayerCheck(contract.name, records, unlisted)


def validate_topology(raw_dir: Path, contract: TopologyContract) -> int:
    """Count topology lines and check the minimum record count."""
    path = _require_file(raw_dir, contract.path, contract.name)
    with path.open("rb") as handle:
        lines = sum(1 for line in handle if line.strip())
    if lines < contract.min_records:
        raise ContractError(f"{contract.path}: has {lines} lines, expected at least {contract.min_records}")
    return lines


@dataclass(frozen=True)
class InputReport:
    layers: dict[str, LayerCheck]
    topology_lines: int
    workbook_rows: dict[str, int]


def validate_inputs(raw_dir: Path, config_dir: Path | None = None) -> InputReport:
    """Validate every declared input. Stops at the first violation with a ContractError."""
    layers = {}
    for name, contract in load_layer_contracts(config_dir).items():
        layers[name] = validate_layer(raw_dir, contract)
    topology_lines = validate_topology(raw_dir, load_topology_contract(config_dir))
    workbook_rows = validate_workbook(raw_dir, load_workbook_contract(config_dir))
    return InputReport(layers, topology_lines, workbook_rows)


def validate_workbook(raw_dir: Path, contract: WorkbookContract) -> dict[str, int]:
    """Check the labor workbook has the expected sheets and enough rows."""
    import openpyxl

    path = _require_file(raw_dir, contract.path, contract.name)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        counts: dict[str, int] = {}
        for sheet in contract.sheets:
            if sheet not in workbook.sheetnames:
                raise ContractError(f"{contract.path}: sheet '{sheet}' is missing")
            counts[sheet] = sum(
                1 for row in workbook[sheet].iter_rows(values_only=True) if any(v is not None for v in row)
            )
        for sheet, minimum in contract.min_rows.items():
            if counts.get(sheet, 0) < minimum:
                raise ContractError(
                    f"{contract.path}: sheet '{sheet}' has {counts.get(sheet, 0)} rows, expected at least {minimum}"
                )
        return counts
    finally:
        workbook.close()
