"""Labor reference: workbook reader, damage-to-labor mapping, size selection and assumptions.

Hours keep the precision stored in the workbook. A labor item with no hour value is stored as
missing, never as zero. Where no direct source exists, the mapping says so and cites an
assumption id.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pandas as pd
import yaml

from .contracts import default_config_dir
from .errors import ConfigError, ContractError

SHEET_CNC = "CNC"
SHEET_MATERIAL = "Labor time material related"
SHEET_LABOR_ONLY = "Labor only work"
SHEET_ILM = "Indirect Labor Multiplier"
SHEET_GROUPS = "CU Group descriptions"

BASES = ("direct", "approximate", "assumed", "none")
UNKNOWN_ASSET = "(unknown)"

# Section title (start of text) -> (group code, how the code was obtained, note)
SECTION_GROUPS: list[tuple[str, str, str, str | None]] = [
    ("Crossarm", "ARM", "workbook", None),
    ("Bracket", "BKT", "workbook", None),
    ("Pole", "POL", "workbook", None),
    ("Transformer", "XFM", "assigned", "The workbook has no group code for transformers."),
    ("Anchor", "ANC", "workbook", None),
    ("Equipment Lead", "EQL", "workbook", None),
    ("Grounding", "GND", "workbook", None),
    ("Guying", "GYF", "workbook", None),
    (
        "Dead-End",
        "DE",
        "assigned",
        "The group descriptions list DEC and DEG separately; this section is not mapped to either.",
    ),
    ("Insulator", "INS", "workbook", None),
    ("Pin", "PIN", "workbook", None),
    ("Recloser", "REC", "workbook", None),
    ("Secondary Conductor Attachment", "SAA", "workbook", None),
    ("Service", "SVC", "workbook", None),
    ("Switch", "SWI", "workbook", None),
    ("Fuse Cutout", "XCO", "workbook", None),
]

ITEM_COLUMNS = [
    "item_id",
    "group_code",
    "group_code_basis",
    "item_key",
    "label",
    "description",
    "size_detail",
    "sample_count",
    "install_hours",
    "remove_hours",
    "install_missing",
    "remove_missing",
    "unit_basis",
    "source_sheet",
    "source_row",
]

INDIRECT_NOTE = (
    "The workbook provides one combined multiplier covering drive time, safety discussion, "
    "setup and similar time. No separate travel or non-productive values were provided."
)


# --------------------------------------------------------------------------- workbook reader
def _number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _item(**kwargs) -> dict:
    row = {column: None for column in ITEM_COLUMNS}
    row.update(kwargs)
    return row


def _read_cnc(ws) -> list[dict]:
    rows = []
    for number, (code, install, remove) in enumerate(ws.iter_rows(min_row=2, max_col=3, values_only=True), start=2):
        if code is None:
            break
        install_h, remove_h = _number(install), _number(remove)
        rows.append(
            _item(
                item_id=f"CNC:{code}",
                group_code="CNC",
                group_code_basis="assigned",
                item_key=str(code),
                description="Connector compatible unit",
                install_hours=install_h,
                remove_hours=remove_h,
                install_missing=install_h is None,
                remove_missing=remove_h is None,
                unit_basis="per unit",
                source_sheet=SHEET_CNC,
                source_row=number,
            )
        )
    return rows


def _section_group(title: str) -> tuple[str, str, str | None]:
    for prefix, code, basis, note in SECTION_GROUPS:
        if title.startswith(prefix):
            return code, basis, note
    raise ContractError(f"{SHEET_MATERIAL}: section '{title}' has no known labor group.")


def _read_material(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    items: list[dict] = []
    index = 0
    while index < len(rows):
        row = rows[index]
        title = row[0]
        is_title = isinstance(title, str) and all(v is None for v in row[1:4])
        if not is_title:
            index += 1
            continue
        group, basis, note = _section_group(title)
        header = [str(v).strip() if v is not None else None for v in rows[index + 1][:4]]
        try:
            install_col = header.index("Avg Install Hours")
            remove_col = header.index("Avg Removal Hours")
        except ValueError as error:
            raise ContractError(f"{SHEET_MATERIAL}: section '{title}' lacks install/removal hour columns.") from error
        cursor = index + 2
        while cursor < len(rows) and any(v is not None for v in rows[cursor][:4]):
            data = rows[cursor]
            if isinstance(data[0], str) and all(v is None for v in data[1:4]) and cursor + 1 < len(rows):
                break  # next section title without a blank row
            fields = dict(zip(header, data[:4]))
            if "Size Group" in fields:
                key, detail = str(fields["Size Group"]), fields.get("Height Range")
            elif "Size (kVA)" in fields:
                size = fields["Size (kVA)"]
                key, detail = f"{int(size)}kVA", f"{int(size)}"
            else:
                key, detail = "average", None
            install_h, remove_h = _number(data[install_col]), _number(data[remove_col])
            count = fields.get("CU Count")
            items.append(
                _item(
                    item_id=f"{group}:{key}",
                    group_code=group,
                    group_code_basis=basis,
                    item_key=key,
                    description=title + (f" ({note})" if note else ""),
                    size_detail=None if detail is None else str(detail),
                    sample_count=None if count is None else int(count),
                    install_hours=install_h,
                    remove_hours=remove_h,
                    install_missing=install_h is None,
                    remove_missing=remove_h is None,
                    unit_basis="per unit",
                    source_sheet=SHEET_MATERIAL,
                    source_row=cursor + 1,
                )
            )
            cursor += 1
        index = cursor
    return items


def _read_labor_only(ws) -> list[dict]:
    items = []
    for number, row in enumerate(ws.iter_rows(min_row=2, max_col=5, values_only=True), start=2):
        name, label, description, hours, unit = row
        if name is None:
            continue
        if description is None:  # later rows hold the description in the second column
            description, label = label, None
        install_h = _number(hours)
        unit_text = str(unit).strip().lower() if unit else ""
        items.append(
            _item(
                item_id=f"ADDER:{name}",
                group_code="ADDER",
                group_code_basis="assigned",
                item_key=str(name),
                label=None if label is None else str(label),
                description=None if description is None else str(description),
                install_hours=install_h,
                remove_hours=None,
                install_missing=install_h is None,
                remove_missing=False,  # the source defines install hours only
                unit_basis=unit_text if unit_text else "per occurrence",
                source_sheet=SHEET_LABOR_ONLY,
                source_row=number,
            )
        )
    return items


def read_workbook(path: Path) -> tuple[pd.DataFrame, dict]:
    """Return (labor_items, parameters) from the workbook. Values only, no formulas evaluated."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        for name in (SHEET_CNC, SHEET_MATERIAL, SHEET_LABOR_ONLY, SHEET_ILM, SHEET_GROUPS):
            if name not in workbook.sheetnames:
                raise ContractError(f"{Path(path).name}: sheet '{name}' is missing")
        items = (
            _read_cnc(workbook[SHEET_CNC])
            + _read_material(workbook[SHEET_MATERIAL])
            + _read_labor_only(workbook[SHEET_LABOR_ONLY])
        )
        ilm_sheet = workbook[SHEET_ILM]
        value = _number(ilm_sheet["B1"].value)
        if value is None or value <= 0:
            raise ContractError(f"{Path(path).name}: the indirect labor multiplier is missing or not positive")
        groups = {}
        for row in workbook[SHEET_GROUPS].iter_rows(min_row=2, max_col=2, values_only=True):
            if row[0]:
                groups[str(row[0]).strip()] = None if row[1] is None else str(row[1]).strip()
    finally:
        workbook.close()

    frame = pd.DataFrame(items, columns=ITEM_COLUMNS)
    frame["sample_count"] = frame["sample_count"].astype("Int64")
    for column in ("install_hours", "remove_hours"):
        frame[column] = frame[column].astype("float64")
    for column in ("install_missing", "remove_missing"):
        frame[column] = frame[column].astype(bool)
    parameters = {
        "indirect_labor_multiplier": {
            "value": value,
            "source_sheet": SHEET_ILM,
            "source_cell": "B1",
            "applies": "Multiply labor hours by this value to include drive time, safety discussion, setup and similar time.",
            "note": INDIRECT_NOTE,
            "overridable": True,
        },
        "workbook_group_descriptions": groups,
        "source_file": Path(path).name,
    }
    return frame, parameters


# --------------------------------------------------------------------------- mapping
@dataclass(frozen=True)
class AssetMapping:
    asset_type: str
    damage_type: str
    labor_group: str | None
    selector: str
    basis: str
    rationale: str
    assumption_id: str | None


@dataclass(frozen=True)
class AdderMapping:
    kind: str  # "access_issue" or "special_equipment"
    value: str
    adder: str | None
    basis: str
    rationale: str
    assumption_id: str | None


@dataclass
class LaborMapping:
    version: int
    review: dict
    assets: dict[tuple[str, str], AssetMapping]
    adders: dict[tuple[str, str], AdderMapping]
    assumptions: dict[str, dict]
    pole_default_group: str = "Medium"


def load_labor_mapping(config_dir: Path | None = None) -> LaborMapping:
    path = (config_dir or default_config_dir()) / "labor_mapping.yaml"
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
    except FileNotFoundError as error:
        raise ConfigError(f"Configuration file not found: {path}") from error
    assets: dict[tuple[str, str], AssetMapping] = {}
    for entry in data["asset_mappings"]:
        mapping = AssetMapping(
            asset_type=str(entry["asset_type"]),
            damage_type=str(entry["damage_type"]),
            labor_group=entry.get("labor_group"),
            selector=entry.get("selector", "none"),
            basis=entry["basis"],
            rationale=entry.get("rationale", ""),
            assumption_id=entry.get("assumption_id"),
        )
        _check_basis(mapping.basis, f"{mapping.asset_type}/{mapping.damage_type}")
        assets[(mapping.asset_type, mapping.damage_type)] = mapping
    adders: dict[tuple[str, str], AdderMapping] = {}
    for kind in ("access_issue", "special_equipment"):
        for value, entry in (data.get("adder_mappings", {}).get(kind) or {}).items():
            mapping = AdderMapping(
                kind=kind,
                value=str(value),
                adder=entry.get("adder"),
                basis=entry["basis"],
                rationale=entry.get("rationale", ""),
                assumption_id=entry.get("assumption_id"),
            )
            _check_basis(mapping.basis, f"{kind}/{value}")
            adders[(kind, mapping.value)] = mapping
    return LaborMapping(
        version=int(data["version"]),
        review=dict(data.get("review") or {}),
        assets=assets,
        adders=adders,
        assumptions={str(k): dict(v) for k, v in (data.get("assumptions") or {}).items()},
    )


def _check_basis(basis: str, where: str) -> None:
    if basis not in BASES:
        raise ConfigError(f"labor_mapping.yaml: {where} has invalid basis '{basis}'")


def validate_mapping(mapping: LaborMapping, items: pd.DataFrame) -> None:
    """Every referenced group and adder must exist and every non-direct row must cite an assumption."""
    groups = set(items["group_code"])
    adder_names = set(items.loc[items["group_code"] == "ADDER", "item_key"])
    for key, entry in mapping.assets.items():
        if entry.labor_group is not None and entry.labor_group not in groups:
            raise ContractError(f"labor_mapping.yaml: {key} points to unknown labor group '{entry.labor_group}'")
        if entry.basis == "none" and entry.labor_group is not None:
            raise ContractError(f"labor_mapping.yaml: {key} has basis 'none' but names a labor group")
        if entry.basis != "direct" and not entry.assumption_id:
            raise ContractError(f"labor_mapping.yaml: {key} is {entry.basis} but cites no assumption id")
    for key, entry in mapping.adders.items():
        if entry.adder is not None and entry.adder not in adder_names:
            raise ContractError(f"labor_mapping.yaml: {key} points to unknown adder '{entry.adder}'")
        if entry.adder is None and entry.basis != "none":
            raise ContractError(f"labor_mapping.yaml: {key} names no adder but basis is '{entry.basis}'")
        if entry.basis != "direct" and not entry.assumption_id:
            raise ContractError(f"labor_mapping.yaml: {key} is {entry.basis} but cites no assumption id")
    cited = {e.assumption_id for e in mapping.assets.values()} | {e.assumption_id for e in mapping.adders.values()}
    for assumption_id in cited - {None}:
        if assumption_id not in mapping.assumptions:
            raise ContractError(f"labor_mapping.yaml: assumption '{assumption_id}' is cited but not defined")


def check_mapping_covers(
    mapping: LaborMapping,
    asset_damage_pairs: set[tuple[str, str]],
    special_equipment: set[str],
    access_values: set[str],
) -> None:
    """Fail when a type found in the damage data has no mapping entry."""
    missing_assets = sorted(p for p in asset_damage_pairs if p not in mapping.assets)
    if missing_assets:
        raise ContractError(f"labor_mapping.yaml has no entry for asset/damage type(s): {missing_assets}")
    missing_equipment = sorted(v for v in special_equipment if ("special_equipment", v) not in mapping.adders)
    if missing_equipment:
        raise ContractError(f"labor_mapping.yaml has no adder entry for special equipment value(s): {missing_equipment}")
    missing_access = sorted(v for v in access_values if v == "Yes" and ("access_issue", v) not in mapping.adders)
    if missing_access:
        raise ContractError("labor_mapping.yaml has no adder entry for an access issue")


# --------------------------------------------------------------------------- selection
@dataclass
class LaborSelection:
    asset_type: str
    damage_type: str
    labor_group: str | None
    item_id: str | None
    install_hours: float | None
    remove_hours: float | None
    basis: str
    assumed: bool
    assumption_ids: list[str] = field(default_factory=list)
    missing: bool = False
    note: str = ""


_RANGE = re.compile(r"(\d+)\D+(\d+)")
_AT_LEAST = re.compile(r"(\d+)\s*ft\s+and\s+above", re.IGNORECASE)


def _height_ranges(rows: pd.DataFrame) -> list[tuple[str, float, float, pd.Series]]:
    ranges = []
    for _, row in rows.iterrows():
        text = str(row["size_detail"])
        above = _AT_LEAST.search(text)
        if above:
            ranges.append((row["item_key"], float(above.group(1)), float("inf"), row))
            continue
        span = _RANGE.search(text)
        if span:
            ranges.append((row["item_key"], float(span.group(1)), float(span.group(2)), row))
    return ranges


def parse_pole_height(value: object) -> float | None:
    if value is None:
        return None
    match = re.search(r"\d+(\.\d+)?", str(value))
    return float(match.group(0)) if match else None


def _row_to_selection(base: dict, row: pd.Series) -> LaborSelection:
    install = None if pd.isna(row["install_hours"]) else float(row["install_hours"])
    remove = None if pd.isna(row["remove_hours"]) else float(row["remove_hours"])
    return LaborSelection(
        item_id=row["item_id"],
        install_hours=install,
        remove_hours=remove,
        missing=install is None and remove is None,
        **base,
    )


def select_labor(
    items: pd.DataFrame,
    mapping: LaborMapping,
    asset_type: str | None,
    damage_type: str | None,
    attributes: dict | None = None,
) -> LaborSelection:
    """Pick labor hours for one damage record from the mapping and the asset's attributes."""
    attributes = attributes or {}
    asset = asset_type or UNKNOWN_ASSET
    damage = damage_type or ""
    entry = mapping.assets.get((asset, damage))
    if entry is None:
        return LaborSelection(
            asset_type=asset, damage_type=damage, labor_group=None, item_id=None, install_hours=None,
            remove_hours=None, basis="none", assumed=False, assumption_ids=[], missing=True,
            note="No mapping entry for this asset and damage type.",
        )
    ids = [entry.assumption_id] if entry.assumption_id else []
    base = dict(asset_type=asset, damage_type=damage, labor_group=entry.labor_group, basis=entry.basis,
                assumed=False, assumption_ids=ids)
    if entry.labor_group is None:
        return LaborSelection(item_id=None, install_hours=None, remove_hours=None, missing=True,
                              note=entry.rationale, **base)

    rows = items[items["group_code"] == entry.labor_group]
    if rows.empty:
        raise ContractError(f"No labor items for group '{entry.labor_group}'")

    if entry.selector == "pole_height":
        ranges = _height_ranges(rows)
        height = parse_pole_height(attributes.get("height"))
        default = rows[rows["item_key"] == mapping.pole_default_group]
        if height is None:
            base.update(assumed=True, assumption_ids=ids + ["A-POLE-01"])
            return _row_to_selection({**base, "note": "Pole height unknown; medium size group assumed."}, default.iloc[0])
        for key, low, high, row in ranges:
            if low <= height <= high:
                return _row_to_selection({**base, "note": f"Size group {key} from pole height."}, row)
        nearest = min(ranges, key=lambda r: min(abs(height - r[1]), abs(height - (r[2] if r[2] != float("inf") else r[1]))))
        base.update(assumed=True, assumption_ids=ids + ["A-POLE-02"])
        return _row_to_selection({**base, "note": f"Height {height:g} ft is outside the documented ranges; nearest size group used."}, nearest[3])

    if entry.selector == "transformer_kva":
        kva = attributes.get("kva")
        if kva is not None and not pd.isna(kva):
            sizes = rows.assign(_kva=rows["size_detail"].astype(float))
            best = sizes.iloc[(sizes["_kva"] - float(kva)).abs().argsort().iloc[0]]
            exact = float(best["_kva"]) == float(kva)
            if not exact:
                base.update(assumed=True, assumption_ids=ids + ["A-XFMR-01"])
            return _row_to_selection({**base, "note": "Size from kVA rating." if exact else "Nearest kVA size used."}, best)
        base.update(assumed=True, assumption_ids=ids + ["A-XFMR-01"])
        mean_row = pd.Series(
            {
                "item_id": f"{entry.labor_group}:mean",
                "install_hours": rows["install_hours"].mean(),
                "remove_hours": rows["remove_hours"].mean(),
            }
        )
        return _row_to_selection({**base, "note": "No kVA rating; unweighted average across documented sizes used."}, mean_row)

    if len(rows) != 1:
        raise ContractError(f"Group '{entry.labor_group}' has {len(rows)} items but selector '{entry.selector}' expects one")
    return _row_to_selection({**base, "note": entry.rationale}, rows.iloc[0])


@dataclass
class AdderSelection:
    kind: str
    value: str
    adder: str | None
    install_hours: float | None
    unit_basis: str | None
    basis: str
    assumption_ids: list[str]
    note: str


def select_adders(
    items: pd.DataFrame, mapping: LaborMapping, access_issue: str | None, special_equipment: str | None
) -> list[AdderSelection]:
    """Adders implied by the access-issue and special-equipment indicators."""
    selections = []
    requested = []
    if access_issue == "Yes":
        requested.append(("access_issue", "Yes"))
    if special_equipment:
        requested.append(("special_equipment", special_equipment))
    for kind, value in requested:
        entry = mapping.adders.get((kind, value))
        if entry is None:
            raise ContractError(f"labor_mapping.yaml has no adder entry for {kind} '{value}'")
        hours, unit = None, None
        if entry.adder is not None:
            row = items[(items["group_code"] == "ADDER") & (items["item_key"] == entry.adder)].iloc[0]
            hours = None if pd.isna(row["install_hours"]) else float(row["install_hours"])
            unit = row["unit_basis"]
        selections.append(
            AdderSelection(kind, value, entry.adder, hours, unit, entry.basis,
                           [entry.assumption_id] if entry.assumption_id else [], entry.rationale)
        )
    return selections


# --------------------------------------------------------------------------- register and output
def build_assumptions_register(mapping: LaborMapping) -> list[dict]:
    """Every assumption cited by a non-direct mapping, plus selection-rule assumptions, with what it affects."""
    affects: dict[str, list[str]] = {}
    for (asset, damage), entry in mapping.assets.items():
        if entry.assumption_id:
            affects.setdefault(entry.assumption_id, []).append(f"asset: {asset} / {damage}")
    for (kind, value), entry in mapping.adders.items():
        if entry.assumption_id:
            affects.setdefault(entry.assumption_id, []).append(f"{kind.replace('_', ' ')}: {value}")
    register = []
    for assumption_id in sorted(mapping.assumptions):
        entry = mapping.assumptions[assumption_id]
        register.append(
            {
                "id": assumption_id,
                "description": entry["description"],
                "affects": sorted(affects.get(assumption_id, entry.get("affects") or [])),
            }
        )
    return register


def mapping_as_output(mapping: LaborMapping) -> dict:
    return {
        "version": mapping.version,
        "review": mapping.review,
        "asset_mappings": [
            {
                "asset_type": e.asset_type,
                "damage_type": e.damage_type,
                "labor_group": e.labor_group,
                "selector": e.selector,
                "basis": e.basis,
                "rationale": e.rationale,
                "assumption_id": e.assumption_id,
            }
            for _, e in sorted(mapping.assets.items())
        ],
        "adder_mappings": [
            {
                "kind": e.kind,
                "value": e.value,
                "adder": e.adder,
                "basis": e.basis,
                "rationale": e.rationale,
                "assumption_id": e.assumption_id,
            }
            for _, e in sorted(mapping.adders.items())
        ],
    }
