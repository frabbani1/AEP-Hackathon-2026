import dataclasses
import shutil

import pytest
import yaml

from aep_pipeline.cli import main
from aep_pipeline.contracts import (
    SENSITIVITIES,
    load_layer_contracts,
    load_topology_contract,
    load_workbook_contract,
    validate_inputs,
    validate_layer,
    validate_topology,
    validate_workbook,
)
from aep_pipeline.errors import ConfigError, ContractError

KEY = {"AEP_TOKEN_KEY": "unit-test-key-0123456789abcdef"}


# ---------------------------------------------------------------- contract files (task 2.2)
def test_all_contract_fields_are_classified():
    contracts = load_layer_contracts()
    assert len(contracts) == 12
    for contract in contracts.values():
        assert contract.fields, contract.name
        for spec in contract.fields.values():
            assert spec.sensitivity in SENSITIVITIES, (contract.name, spec.name)
            assert spec.description, (contract.name, spec.name)


def test_sensitive_customer_fields_are_never_public_or_internal_outputs():
    customer = load_layer_contracts()["Customer_Clip"]
    for name in ("critical_c", "address", "address_nu", "road_name", "street_cod", "city", "macss_mete",
                 "anno_strin", "service_po", "service__1", "special_op", "aep_locati"):
        assert customer.fields[name].sensitivity == "drop", name
    assert customer.fields["esri_defau"].sensitivity == "undocumented"
    assert customer.fields["premise_nu"].sensitivity == "internal"
    assert set(customer.columns()) == {"premise_nu", "circuit_id", "aep_xfmr_l", "customer_c", "kind_of_se"}


def test_priority_layers_only_keep_category_codes():
    for name in ("Ambulance_Servi_Pairwis_Clip", "Fire_Stations_E_Pairwis_Clip", "Hospitals_Medic_Pairwis_Clip"):
        contract = load_layer_contracts()[name]
        assert set(contract.columns()) == {"FTYPE", "FCODE"}
        assert contract.fields["NAME"].sensitivity == "drop"
        assert contract.fields["ADDRESS"].sensitivity == "drop"


def test_control_ids_global_ids_and_crew_fields_are_not_readable():
    contracts = load_layer_contracts()
    for layer in ("Fuse_Bank_Clip", "Transformer_Clip", "Recloser_Bank_Clip", "Switch_Clip"):
        assert contracts[layer].fields["scada_cont"].sensitivity == "drop"
        assert contracts[layer].fields["global_id"].sensitivity == "drop"
    arcos = contracts["arcos_prod_damage_09_28_Clip"]
    for name in ("TASK_Crew", "TASK_Job_I", "Crew_Area", "Reporting_", "Reporting1", "Repaired_B"):
        assert arcos.fields[name].sensitivity == "drop", name


def test_topology_and_workbook_contracts_load():
    assert load_topology_contract().node_parts == ("object_type", "object_id", "subtype", "device_key")
    assert "Indirect Labor Multiplier" in load_workbook_contract().sheets


@pytest.mark.real_data
def test_real_inputs_satisfy_their_contracts(real_data_dir):
    report = validate_inputs(real_data_dir)
    assert report.layers["Customer_Clip"].records == 117642
    assert report.topology_lines == 117599
    assert not any(check.unlisted_fields for check in report.layers.values())


# ---------------------------------------------------------------- validator (task 2.3)
def test_fixture_inputs_validate(fixture_raw, test_config_dir):
    report = validate_inputs(fixture_raw.raw_dir, test_config_dir)
    assert report.layers["Transformer_Clip"].records == 9


def test_missing_field_names_the_file_and_field(fixture_raw, test_config_dir):
    contract = load_layer_contracts(test_config_dir)["Fuse_Bank_Clip"]
    extra = {**contract.fields, "bogus_field": dataclasses.replace(contract.fields["id"], name="bogus_field")}
    broken = dataclasses.replace(contract, fields=extra)
    with pytest.raises(ContractError) as error:
        validate_layer(fixture_raw.raw_dir, broken)
    assert "Data1/Fuse_Bank_Clip.shp" in str(error.value)
    assert "bogus_field" in str(error.value)
    assert error.value.exit_code == 1


def test_wrong_geometry_type(fixture_raw, test_config_dir):
    contract = load_layer_contracts(test_config_dir)["Primary_ExportFeatures"]
    broken = dataclasses.replace(contract, geometry_type="Point")
    with pytest.raises(ContractError) as error:
        validate_layer(fixture_raw.raw_dir, broken)
    assert "Primary_ExportFeatures.shp" in str(error.value)
    assert "geometry type" in str(error.value)
    assert "LineString" in str(error.value)


def test_wrong_field_type(fixture_raw, test_config_dir):
    contract = load_layer_contracts(test_config_dir)["Pole_Clip"]
    fields = dict(contract.fields)
    fields["id"] = dataclasses.replace(fields["id"], dtype="float64")
    with pytest.raises(ContractError) as error:
        validate_layer(fixture_raw.raw_dir, dataclasses.replace(contract, fields=fields))
    assert "field 'id'" in str(error.value)
    assert "float64" in str(error.value)


def test_record_count_below_minimum_reports_both_counts(fixture_raw, test_config_dir):
    contract = load_layer_contracts(test_config_dir)["Switch_Clip"]
    broken = dataclasses.replace(contract, min_records=1000)
    with pytest.raises(ContractError) as error:
        validate_layer(fixture_raw.raw_dir, broken)
    message = str(error.value)
    assert "has 2 records" in message and "at least 1000" in message


def test_wrong_reference_system(fixture_raw, test_config_dir):
    contract = load_layer_contracts(test_config_dir)["Switch_Clip"]
    broken = dataclasses.replace(contract, crs="EPSG:3857")
    with pytest.raises(ContractError) as error:
        validate_layer(fixture_raw.raw_dir, broken)
    assert "reference system" in str(error.value)


def test_git_lfs_pointer_is_detected_with_a_hint(tmp_path, fixture_raw, test_config_dir):
    raw = tmp_path / "data"
    shutil.copytree(fixture_raw.raw_dir, raw)
    pointer = b"version https://git-lfs.github.com/spec/v1\noid sha256:abc\nsize 123\n"
    (raw / "Data1" / "Customer_Clip.dbf").write_bytes(pointer)
    contract = load_layer_contracts(test_config_dir)["Customer_Clip"]
    with pytest.raises(ContractError) as error:
        validate_layer(raw, contract)
    assert "Customer_Clip.dbf" in str(error.value)
    assert "git lfs pull" in str(error.value)


def test_missing_file_is_a_contract_error(tmp_path, test_config_dir):
    contract = load_layer_contracts(test_config_dir)["Pole_Clip"]
    with pytest.raises(ContractError) as error:
        validate_layer(tmp_path, contract)
    assert "Data1/Pole_Clip.shp" in str(error.value)


def test_topology_too_short(fixture_raw, test_config_dir):
    contract = dataclasses.replace(load_topology_contract(test_config_dir), min_records=500)
    with pytest.raises(ContractError) as error:
        validate_topology(fixture_raw.raw_dir, contract)
    assert "matched_upstream.txt" in str(error.value)


def test_workbook_missing_sheet(tmp_path, fixture_raw, test_config_dir):
    import openpyxl

    raw = tmp_path / "data"
    shutil.copytree(fixture_raw.raw_dir, raw)
    path = raw / "Install_Labor_HackOHIO.xlsx"
    wb = openpyxl.load_workbook(path)
    del wb["Indirect Labor Multiplier"]
    wb.save(path)
    with pytest.raises(ContractError) as error:
        validate_workbook(raw, load_workbook_contract(test_config_dir))
    assert "Indirect Labor Multiplier" in str(error.value)


def test_cli_exits_1_on_a_contract_violation(tmp_path, fixture_raw, test_config_dir, capsys):
    config = tmp_path / "config"
    shutil.copytree(test_config_dir, config)
    path = config / "contracts" / "Customer_Clip.yaml"
    data = yaml.safe_load(path.read_text())
    data["fields"]["not_in_the_layer"] = {"type": "int64", "sensitivity": "public", "description": "test"}
    path.write_text(yaml.safe_dump(data, sort_keys=False))
    code = main(
        ["run", "--raw", str(fixture_raw.raw_dir), "--out", str(tmp_path / "out"), "--config", str(config)],
        environ=KEY,
    )
    err = capsys.readouterr().err
    assert code == 1
    assert "Customer_Clip.shp" in err and "not_in_the_layer" in err


def test_invalid_sensitivity_in_a_contract_is_a_configuration_error(tmp_path, test_config_dir):
    config = tmp_path / "config"
    shutil.copytree(test_config_dir, config)
    path = config / "contracts" / "Customer_Clip.yaml"
    data = yaml.safe_load(path.read_text())
    data["fields"]["customer_c"]["sensitivity"] = "maybe"
    path.write_text(yaml.safe_dump(data, sort_keys=False))
    with pytest.raises(ConfigError):
        load_layer_contracts(config)
