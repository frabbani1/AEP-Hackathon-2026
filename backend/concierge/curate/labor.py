"""Parse Install_Labor_HackOHIO.xlsx into tidy labor standards (units, adders, indirect multiplier)."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

SOURCE = "Install_Labor_HackOHIO.xlsx"
REQUIRED_CATEGORIES = {"pole", "transformer", "recloser", "switch", "fuse_cutout", "crossarm", "connector"}


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def _category(title: str) -> str:
    t = re.sub(r"\s+(CU\s+)?Average Hours.*$", "", title.strip(), flags=re.I)
    return slug(t)


def _num(v) -> float | None:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _parse_blocks(rows: list[tuple]) -> list[dict]:
    """Blocks look like: title row, header row, then one or more numeric rows, then a blank row."""
    units: list[dict] = []
    i = 0
    while i < len(rows):
        title = rows[i][0]
        if isinstance(title, str) and "Average Hours" in title:
            cat = _category(title)
            header = [h if isinstance(h, str) else None for h in rows[i + 1]]
            idx = {h: n for n, h in enumerate(header) if h}
            j = i + 2
            while j < len(rows) and any(c is not None for c in rows[j]):
                row = rows[j]
                inst = _num(row[idx["Avg Install Hours"]]) if "Avg Install Hours" in idx else None
                rem = _num(row[idx["Avg Removal Hours"]]) if "Avg Removal Hours" in idx else None
                if inst is not None or rem is not None:
                    size_label = size_kva = height = None
                    if "Size Group" in idx:
                        size_label, height = str(row[idx["Size Group"]]), str(row[idx["Height Range"]])
                    if "Size (kVA)" in idx:
                        size_kva = int(row[idx["Size (kVA)"]])
                    units.append({
                        "category": cat, "size_group": size_label, "height_range": height, "size_kva": size_kva,
                        "cu_count": int(row[idx["CU Count"]]) if "CU Count" in idx and row[idx["CU Count"]] else None,
                        "install_hours": inst, "remove_hours": rem,
                    })
                j += 1
            i = j
        else:
            i += 1
    return units


def parse_labor(path: str | Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)

    blocks = [tuple(r) for r in wb["Labor time material related"].iter_rows(values_only=True)]
    units = _parse_blocks(blocks)

    # Connector average sits to the right of the long CNC list (E1:G3)
    cnc = [tuple(r) for r in wb["CNC"].iter_rows(min_row=1, max_row=3, min_col=5, max_col=7, values_only=True)]
    units.append({"category": "connector", "size_group": None, "height_range": None, "size_kva": None,
                  "cu_count": int(cnc[2][0]), "install_hours": _num(cnc[2][1]), "remove_hours": _num(cnc[2][2])})

    # Rows 2-8: (original intent, anonymized work type, description, hours, unit); we keep the anonymized name.
    # Rows 9-12: (adder name, "apply when ..." text, -, hours). Column C is empty, so B is the description.
    adders = []
    for r in wb["Labor only work"].iter_rows(min_row=2, values_only=True):
        if not any(c is not None for c in r):
            continue
        a_, b_, c_, d_, e_ = (list(r) + [None] * 5)[:5]
        if c_ is not None:
            key, description = slug(b_), c_
        else:
            key, description = slug(a_), b_
        adders.append({"key": key, "description": description, "hours": _num(d_), "unit": e_ or "per job"})

    ilm = _num(wb["Indirect Labor Multiplier"]["B1"].value)
    groups = [{"code": r[0], "description": r[1]}
              for r in wb["CU Group descriptions"].iter_rows(min_row=2, values_only=True) if r[0]]

    out = {"source": SOURCE, "indirect_labor_multiplier": ilm, "units": units, "adders": adders, "cu_groups": groups}
    validate_labor(out)
    return out


def validate_labor(d: dict) -> None:
    cats = {u["category"] for u in d["units"]}
    missing = REQUIRED_CATEGORIES - cats
    if missing:
        raise ValueError(f"labor workbook is missing categories: {sorted(missing)}")
    if d["indirect_labor_multiplier"] is None or d["indirect_labor_multiplier"] <= 1:
        raise ValueError("indirect labor multiplier not found or implausible")
    if any(u["install_hours"] is not None and u["install_hours"] < 0 for u in d["units"]):
        raise ValueError("negative labor hours")
